"""文件解析服务 — Sprint 2.A。

输入:磁盘文件路径 + mime_type
输出:纯文本 str(用于:① 红旗扫描 ② 后续 2.B 抽图谱)

支持格式:
  text/plain                                                 .txt    直读 utf-8(失败 fallback gbk)
  application/epub+zip                                       .epub   ebooklib + bs4 抽 plaintext
  application/vnd.openxmlformats-officedocument.wordprocessingml.document .docx   python-docx 抽段落

设计:
  - 不抛异常给上层,统一返回 ParseResult(成功 / 失败原因)
  - epub / docx 内嵌恶意脚本不会被执行(我们只抽文本字段)
  - 大文件解析无流式优化(50-100MB 量级,内存够用)
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional


SUPPORTED_MIMES = {
    "text/plain",
    "application/epub+zip",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

EXT_TO_MIME = {
    ".txt": "text/plain",
    ".epub": "application/epub+zip",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass
class ParseResult:
    success: bool
    text: str = ""
    error: Optional[str] = None

    @classmethod
    def ok(cls, text: str) -> "ParseResult":
        return cls(success=True, text=text)

    @classmethod
    def fail(cls, error: str) -> "ParseResult":
        return cls(success=False, error=error)


def detect_ext(filename: str) -> Optional[str]:
    """从文件名末尾抽小写扩展名。无扩展名 → None。"""
    name = filename.lower()
    for ext in EXT_TO_MIME:
        if name.endswith(ext):
            return ext
    return None


def mime_matches_ext(mime: str, ext: Optional[str]) -> bool:
    """mime 与扩展名双重校验,防伪造扩展名上传可执行文件。"""
    if not ext:
        return False
    expected = EXT_TO_MIME.get(ext)
    return expected is not None and mime == expected


def parse_file(path: Path, mime: str) -> ParseResult:
    """主入口。"""
    if mime not in SUPPORTED_MIMES:
        return ParseResult.fail(f"不支持的文件类型:{mime}")
    try:
        if mime == "text/plain":
            return _parse_txt(path)
        if mime == "application/epub+zip":
            return _parse_epub(path)
        if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            return _parse_docx(path)
    except Exception as e:  # noqa: BLE001
        return ParseResult.fail(f"解析异常({type(e).__name__}):{e}")
    return ParseResult.fail(f"未实现的解析路径:{mime}")


# === txt ===

def _parse_txt(path: Path) -> ParseResult:
    raw = path.read_bytes()
    # 优先 utf-8(BOM 也吃),失败 fallback gbk(常见中文 Windows 文件)
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            text = raw.decode(encoding)
            return ParseResult.ok(_normalize(text))
        except UnicodeDecodeError:
            continue
    return ParseResult.fail("无法识别文本编码(尝试了 utf-8 / gbk)")


# === epub ===

def _parse_epub(path: Path) -> ParseResult:
    """ebooklib 读 epub → 按 spine 顺序抽各章 HTML 的 plaintext。

    2026-05-28 修(治"解析出 0 字"):原用 get_items_of_type(ITEM_DOCUMENT) 抽取,
    但部分 epub(扫描 OCR 版 / 非标准 media-type)正文 item 被 ebooklib 归为
    ITEM_UNKNOWN(type 0),get_items_of_type(ITEM_DOCUMENT) 只拿到空的 nav.xhtml → 0 字。
    改为按 spine(epub 规范里"正文阅读顺序"的权威来源)遍历,覆盖所有正文 item 类型;
    再用 ITEM_DOCUMENT 兜底 spine 没覆盖到的(极少数 spine 不全的 epub)。
    """
    from ebooklib import epub, ITEM_DOCUMENT
    from bs4 import BeautifulSoup

    book = epub.read_epub(str(path), options={"ignore_ncx": True})

    _SKIP_SUFFIXES = (".css", ".ncx", ".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".otf", ".ttf")

    def _extract(content: bytes) -> str:
        soup = BeautifulSoup(content, "html.parser")
        # script / style 抽掉(html.parser 不执行脚本,但避免噪声文本)
        for tag in soup(["script", "style"]):
            tag.decompose()
        return soup.get_text(separator="\n").strip()

    chunks: list[str] = []
    seen_ids: set[str] = set()

    # 1. 优先按 spine 顺序(正文阅读顺序权威来源,兼容非标准 media-type 的正文 item)
    for idref, _ in book.spine:
        item = book.get_item_with_id(idref)
        if item is None:
            continue
        seen_ids.add(item.get_id())
        name = (item.get_name() or "").lower()
        if name.endswith(_SKIP_SUFFIXES):
            continue
        text = _extract(item.get_content())
        if text:
            chunks.append(text)

    # 2. 兜底:spine 没覆盖到的 ITEM_DOCUMENT(标准 epub 通常已被 spine 覆盖)
    for item in book.get_items_of_type(ITEM_DOCUMENT):
        if item.get_id() in seen_ids:
            continue
        text = _extract(item.get_content())
        if text:
            chunks.append(text)

    if not chunks:
        return ParseResult.fail(
            "epub 内未找到任何文本内容(可能是纯图片扫描版,无文本层)"
        )
    return ParseResult.ok(_normalize("\n\n".join(chunks)))


# === docx ===

def _parse_docx(path: Path) -> ParseResult:
    """python-docx 抽段落 + 表格单元格文本。"""
    from docx import Document

    doc = Document(str(path))
    chunks: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            chunks.append(para.text)
    # 表格也抽(部分小说类排版会用表格)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    chunks.append(cell.text)
    if not chunks:
        return ParseResult.fail("docx 内未找到任何段落内容")
    return ParseResult.ok(_normalize("\n\n".join(chunks)))


# === 文本规范化 ===

def _normalize(text: str) -> str:
    """规范化:统一换行符 + 折叠 OCR 中文字间空格 + 折叠 3+ 空行为 2 个。"""
    import re

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # OCR 扫描件(如 Internet Archive 转的 epub)常把中文拆成单字带空格
    # ("给 那 贫 子")→ 删汉字与汉字之间的行内空格(空格 / 制表 / 全角空格).
    # 设计:
    #   - lookbehind/lookahead 都要求汉字 → 对英文/数字间空格(hello world / 第 3 章)无影响
    #   - 字符类不含 \n → 保留段落换行结构
    #   - 对正常中文(汉字间本无空格)是 no-op,无副作用
    text = re.sub(
        r"(?<=[一-鿿])[ \t　]+(?=[一-鿿])", "", text
    )
    # 折叠超过 2 个的连续空行(避免 epub 转换出大量空白)
    while "\n\n\n" in text:
        text = text.replace("\n\n\n", "\n\n")
    return text.strip()


__all__ = [
    "SUPPORTED_MIMES",
    "EXT_TO_MIME",
    "ParseResult",
    "parse_file",
    "detect_ext",
    "mime_matches_ext",
]
